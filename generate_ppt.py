import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    BG_DARK = RGBColor(11, 19, 43)        # #0b132b
    BG_LIGHT = RGBColor(248, 250, 252)    # #f8fafc
    WHITE = RGBColor(255, 255, 255)
    PRIMARY = RGBColor(0, 82, 204)        # #0052cc
    ACCENT = RGBColor(0, 199, 230)        # #00c7e6
    DARK_TEXT = RGBColor(15, 23, 42)      # #0f172a
    MUTED_TEXT = RGBColor(100, 116, 139)  # #64748b
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240) # #e2e8f0
    SUCCESS = RGBColor(16, 185, 129)      # #10b981
    TAG_BG = RGBColor(238, 242, 255)

    def add_header(slide, title_text, category_text):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = PRIMARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = DARK_TEXT

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.color.rgb = BG_DARK

    # Sub-badge
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(3.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(22, 36, 71)
    badge.line.color.rgb = PRIMARY
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "FINAL PROJECT DEFENSE • 7-MIN PRESENTATION"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT
    p_b.alignment = PP_ALIGN.CENTER

    # Main Title
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "BETE (FixLink)"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf1.add_paragraph()
    p2.text = "On-Demand Home Services & Verified Technician Accreditation Marketplace"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_before = Pt(10)

    # Info bar card
    info_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.7), Inches(11.333), Inches(1.8))
    info_card.fill.solid()
    info_card.fill.fore_color.rgb = RGBColor(17, 28, 56)
    info_card.line.color.rgb = RGBColor(30, 41, 59)
    itf = info_card.text_frame
    itf.word_wrap = True

    ip1 = itf.paragraphs[0]
    ip1.text = "Core Presentation Pillars:"
    ip1.font.size = Pt(13)
    ip1.font.bold = True
    ip1.font.color.rgb = ACCENT

    ip2 = itf.add_paragraph()
    ip2.text = "• 1. End-to-End Customer & Technician Workflow (Mobile App - React Native / Expo)"
    ip2.font.size = Pt(12)
    ip2.font.color.rgb = WHITE
    ip2.space_before = Pt(6)

    ip3 = itf.add_paragraph()
    ip3.text = "• 2. Real-Time Accreditation & Verification Portal (Web Portal - React 19 + Vite)"
    ip3.font.size = Pt(12)
    ip3.font.color.rgb = WHITE

    ip4 = itf.add_paragraph()
    ip4.text = "• 3. Micro-Architecture: Node.js, Express, MongoDB Atlas, Socket.io, Cloudinary & Chapa API"
    ip4.font.size = Pt(12)
    ip4.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 2: Problem & Solution (60-sec intro)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Market Problem & The Bete Solution", "01 / Context & Value Proposition")

    # Left Card: The Problem
    card_prob = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.1))
    card_prob.fill.solid()
    card_prob.fill.fore_color.rgb = RGBColor(254, 242, 242) # light red
    card_prob.line.color.rgb = RGBColor(254, 202, 202)
    ptf = card_prob.text_frame
    ptf.word_wrap = True

    pp0 = ptf.paragraphs[0]
    pp0.text = "THE PROBLEM IN ADDIS ABABA"
    pp0.font.size = Pt(14)
    pp0.font.bold = True
    pp0.font.color.rgb = RGBColor(185, 28, 28)

    problems = [
        "Unvetted Technicians: Finding trustworthy plumbers, electricians, or appliance mechanics is based purely on word-of-mouth with zero accreditation.",
        "Security & Safety Risks: Homeowners hesitate to allow unknown handymen into their homes without verified identity documentation.",
        "No Transparent Bidding: Customers face arbitrary price surges with no counter-offers or structured job scope estimation.",
        "Disconnected Service Seekers: Qualified technicians lack a digital presence to showcase certifications and earn verified credibility."
    ]
    for p in problems:
        para = ptf.add_paragraph()
        para.text = "• " + p
        para.font.size = Pt(12)
        para.font.color.rgb = DARK_TEXT
        para.space_before = Pt(12)

    # Right Card: The Bete Solution
    card_sol = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.1))
    card_sol.fill.solid()
    card_sol.fill.fore_color.rgb = RGBColor(240, 253, 250) # light green/teal
    card_sol.line.color.rgb = RGBColor(153, 246, 228)
    stf = card_sol.text_frame
    stf.word_wrap = True

    sp0 = stf.paragraphs[0]
    sp0.text = "THE BETE (FIXLINK) SOLUTION"
    sp0.font.size = Pt(14)
    sp0.font.bold = True
    sp0.font.color.rgb = RGBColor(13, 148, 136)

    solutions = [
        "Mandatory Credential Verification: Technicians must submit National ID and Trade Certificates for admin review before receiving the Verified Badge.",
        "Geo-Targeted Subcity Matching: Jobs are categorized by subcity (Bole, Yeka, Kirkos, etc.) allowing fast local responses.",
        "Competitive Proposal Marketplace: Technicians submit transparent bids; customers inspect ratings, portfolios, and bids before accepting.",
        "Connects & Escrow Model: Anti-spam system using Connects balance to bid, backed by integrated Chapa payments (Telebirr & CBE)."
    ]
    for s in solutions:
        para = stf.add_paragraph()
        para.text = "✓ " + s
        para.font.size = Pt(12)
        para.font.color.rgb = DARK_TEXT
        para.space_before = Pt(12)

    # ==========================================
    # SLIDE 3: System Architecture
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Full-Stack System Architecture", "02 / Engineering Blueprint")

    cols = [
        ("Mobile Client", "React Native + Expo", "Cross-platform mobile client for iOS & Android.\n• Customer booking flow\n• Technician bidding feed\n• In-app chat interface\n• Document upload UI", PRIMARY),
        ("Web Admin Portal", "React 19 + Vite", "High-performance administrative console.\n• Technician document inspection\n• 1-Click Approve / Reject\n• User management directory\n• Real-time marketplace metrics", RGBColor(14, 116, 144)),
        ("REST & Socket API", "Node.js + Express", "Asynchronous central server.\n• JWT & Bcrypt auth middleware\n• Role-Based Access Control\n• Real-time Socket.io engine\n• Chapa webhook receiver", RGBColor(5, 150, 105)),
        ("Database & Cloud", "MongoDB Atlas + Cloudinary", "Cloud data & media persistence.\n• MongoDB Mongoose schemas\n• Dynamic document attributes\n• Cloudinary CDN image delivery\n• Indexed subcity queries", RGBColor(124, 58, 237))
    ]

    for i, (title, tech, desc, border_col) in enumerate(cols):
        x = Inches(0.8 + i * 2.95)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.8), Inches(4.9))
        c.fill.solid()
        c.fill.fore_color.rgb = WHITE
        c.line.color.rgb = border_col
        c.line.width = Pt(2)
        ctf = c.text_frame
        ctf.word_wrap = True

        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(15)
        cp1.font.bold = True
        cp1.font.color.rgb = DARK_TEXT

        cp2 = ctf.add_paragraph()
        cp2.text = tech
        cp2.font.size = Pt(11)
        cp2.font.bold = True
        cp2.font.color.rgb = border_col
        cp2.space_before = Pt(4)

        cp3 = ctf.add_paragraph()
        cp3.text = desc
        cp3.font.size = Pt(11)
        cp3.font.color.rgb = RGBColor(71, 85, 105)
        cp3.space_before = Pt(14)

    # ==========================================
    # SLIDE 4: "Which Technology Powers Which Action?" (THE KEY DEFENSE SLIDE)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Technology Breakdown by Action & Feature", "03 / Tech Stack Defense Matrix")

    # Create Table
    table_shape = s4.shapes.add_table(7, 4, Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.1))
    table = table_shape.table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(2.6)
    table.columns[2].width = Inches(3.8)
    table.columns[3].width = Inches(2.1)

    headers = ["User Action / Feature", "Technology Used", "Architectural Justification", "Source Component"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE

    rows_data = [
        ("Mobile Interface & Bidding", "React Native + Expo", "Single declarative codebase for Android & iOS with native performance", "mobile/"),
        ("Admin Verification Queue", "React 19 + Vite", "Rapid sub-second HMR, modular views, and clean unstyled inspection layout", "website/src/admin/"),
        ("Real-Time Chat & Updates", "Socket.io (WebSockets)", "Bi-directional persistent room channels with zero polling overhead", "chatSocket.js"),
        ("National ID & Certificate Storage", "Multer + Cloudinary", "Direct CDN upload with automatic image compression and asset resolution", "cloudinary.js"),
        ("Role Access & Passwords", "JWT + Bcrypt (10 rounds)", "Stateless bearer authorization with salted hashing for passwords", "authMiddleware.js"),
        ("Connects & Fee Payments", "Chapa API", "Native support for Ethiopian payment rails (Telebirr, CBE Birr)", "chapaService.js")
    ]

    for row_idx, row in enumerate(rows_data):
        for col_idx, val in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(248, 250, 252) if row_idx % 2 == 0 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = DARK_TEXT
            if col_idx == 1:
                p.font.bold = True
                p.font.color.rgb = PRIMARY

    # ==========================================
    # SLIDE 5: Live Demo Roadmap (The 4-Minute Demo Guide)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Live Demonstration Roadmap (4 Minutes)", "04 / Demonstration Script")

    steps = [
        ("Step 1: Service Request Posting", "Customer Experience (Mobile)", "• Customer selects service category (e.g. Electrical / Appliances)\n• Inputs specific description, selects subcity (e.g. Bole), attaches photo\n• Job is saved in MongoDB and broadcast to eligible technicians", PRIMARY),
        ("Step 2: Competitive Proposal & Chat", "Technician Experience (Mobile)", "• Technician opens job feed, checks subcity distance\n• Submits competitive price quote (deducts 2 Connects from wallet)\n• Customer & technician initiate live Socket.io messaging room", RGBColor(14, 116, 144)),
        ("Step 3: Accreditation Verification", "Admin Verification (Web Console)", "• Admin opens Verification Queue (Approve / Reject tabs)\n• Inspects high-res National ID & Trade Certificate / License\n• 1-Click Approve: Sets verified status & awards 10 welcome Connects", RGBColor(5, 150, 105))
    ]

    for i, (stitle, role, details, color) in enumerate(steps):
        x = Inches(0.8 + i * 3.95)
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.8), Inches(4.9))
        c.fill.solid()
        c.fill.fore_color.rgb = WHITE
        c.line.color.rgb = color
        c.line.width = Pt(2)
        ctf = c.text_frame
        ctf.word_wrap = True

        cp1 = ctf.paragraphs[0]
        cp1.text = stitle
        cp1.font.size = Pt(14)
        cp1.font.bold = True
        cp1.font.color.rgb = DARK_TEXT

        cp2 = ctf.add_paragraph()
        cp2.text = role
        cp2.font.size = Pt(11)
        cp2.font.bold = True
        cp2.font.color.rgb = color
        cp2.space_before = Pt(4)

        cp3 = ctf.add_paragraph()
        cp3.text = details
        cp3.font.size = Pt(11)
        cp3.font.color.rgb = RGBColor(71, 85, 105)
        cp3.space_before = Pt(14)

    # ==========================================
    # SLIDE 6: Security, Real-Time & Cloud Deployment
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Security, Real-Time Architecture & Deployment", "05 / Engineering Rigor")

    sec_cards = [
        ("Authentication & RBAC", "• Passwords hashed with 10 salt rounds of Bcrypt\n• Stateless JWT tokens signed with HMAC SHA-256\n• Granular route guards: authMiddleware, adminMiddleware\n• Protected document routes preventing unauthorized scraping"),
        ("Real-Time Synchronization", "• WebSocket communication powered by Socket.io\n• Job-isolated rooms: socket.join(`job_${jobId}`)\n• Instant event triggers: new_message, bid_submitted\n• Eliminates battery-draining HTTP polling on mobile"),
        ("Cloud Production Deployment", "• Frontend: Deployed on Vercel with automated GitHub CI/CD\n• SPA Routing: Managed via vercel.json rewrites\n• Backend: Hosted on Render Web Service\n• Database: Hosted on MongoDB Atlas with automated cloud backups")
    ]

    for i, (title, points) in enumerate(sec_cards):
        y = Inches(1.8 + i * 1.7)
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(1.45))
        c.fill.solid()
        c.fill.fore_color.rgb = WHITE
        c.line.color.rgb = CARD_BORDER
        ctf = c.text_frame
        ctf.word_wrap = True

        cp1 = ctf.paragraphs[0]
        cp1.text = title
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = PRIMARY

        cp2 = ctf.add_paragraph()
        cp2.text = points
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = DARK_TEXT
        cp2.space_before = Pt(4)

    # ==========================================
    # SLIDE 7: Conclusion & Q&A
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    bg7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg7.fill.solid()
    bg7.fill.fore_color.rgb = BG_DARK
    bg7.line.color.rgb = BG_DARK

    q_box = s7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(2.0))
    qtf = q_box.text_frame
    qtf.word_wrap = True
    qp1 = qtf.paragraphs[0]
    qp1.text = "Thank You! Ready for Q & A"
    qp1.font.size = Pt(40)
    qp1.font.bold = True
    qp1.font.color.rgb = WHITE

    qp2 = qtf.add_paragraph()
    qp2.text = "Bete (FixLink) — Empowering Ethiopian Households & Vetted Service Technicians"
    qp2.font.size = Pt(18)
    qp2.font.color.rgb = ACCENT
    qp2.space_before = Pt(8)

    # Summary card for Q&A reference
    sum_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.0), Inches(11.3), Inches(2.5))
    sum_card.fill.solid()
    sum_card.fill.fore_color.rgb = RGBColor(17, 28, 56)
    sum_card.line.color.rgb = RGBColor(30, 41, 59)
    stf7 = sum_card.text_frame
    stf7.word_wrap = True

    stp1 = stf7.paragraphs[0]
    stp1.text = "Quick Technical Reference for Evaluators:"
    stp1.font.size = Pt(13)
    stp1.font.bold = True
    stp1.font.color.rgb = WHITE

    refs = [
        "• Mobile Frontend: React Native, Expo, React Navigation, Axios",
        "• Admin & Web: React 19, Vite, Lucide Icons, Modern Vanilla CSS",
        "• Backend & Real-Time: Node.js, Express.js, Socket.io, JWT, Multer",
        "• Cloud Infrastructure: MongoDB Atlas, Cloudinary CDN, Chapa Payment Gateway, Vercel & Render"
    ]
    for r in refs:
        rp = stf7.add_paragraph()
        rp.text = r
        rp.font.size = Pt(11)
        rp.font.color.rgb = RGBColor(203, 213, 225)
        rp.space_before = Pt(4)

    output_path = "Bete_7Min_Defense_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_presentation()
